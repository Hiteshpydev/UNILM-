from pptx import Presentation
from pathlib import Path


def parse_pptx(file_path):
    """
    Extract text, tables, and embedded images
    from a PowerPoint presentation.
    """

    try:
        presentation = Presentation(file_path)

        data = []

        slide_count = 0
        table_count = 0
        image_count = 0

        image_dir = Path("extracted_images")
        image_dir.mkdir(exist_ok=True)

        # --------------------------------
        # Process slides
        # --------------------------------

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1
        ):

            slide_count += 1

            slide_text = []

            # --------------------------------
            # 1. Text
            # --------------------------------

            for shape in slide.shapes:

                if hasattr(shape, "text"):

                    text = shape.text.strip()

                    if text:
                        slide_text.append(text)

            combined_text = "\n".join(
                slide_text
            )

            if combined_text:

                data.append({
                    "location": {
                        "type": "slide",
                        "number": slide_number
                    },
                    "content": {
                        "type": "text",
                        "text": combined_text
                    }
                })

            # --------------------------------
            # 2. Tables
            # --------------------------------

            for shape in slide.shapes:

                if not shape.has_table:
                    continue

                table_data = []

                for row in shape.table.rows:

                    row_data = []

                    for cell in row.cells:

                        row_data.append(
                            cell.text.strip()
                        )

                    table_data.append(row_data)

                if not any(
                    any(cell for cell in row)
                    for row in table_data
                ):
                    continue

                table_count += 1

                data.append({
                    "location": {
                        "type": "slide",
                        "number": slide_number
                    },
                    "content": {
                        "type": "table",
                        "data": table_data
                    }
                })

            # --------------------------------
            # 3. Embedded images
            # --------------------------------

            for shape in slide.shapes:

                if shape.shape_type != 13:
                    continue

                image = shape.image

                image_count += 1

                image_ext = image.ext

                image_path = (
                    image_dir
                    / f"{Path(file_path).stem}"
                    f"_slide_{slide_number}"
                    f"_image_{image_count}"
                    f".{image_ext}"
                )

                with open(
                    image_path,
                    "wb"
                ) as image_file:

                    image_file.write(
                        image.blob
                    )

                data.append({
                    "location": {
                        "type": "image",
                        "number": image_count
                    },
                    "content": {
                        "type": "image",
                        "image_path": str(
                            image_path
                        ),
                        "source": "embedded"
                    }
                })

        # --------------------------------
        # Metadata
        # --------------------------------

        metadata = {
            "source": Path(file_path).name,
            "file_type": "pptx",
            "slide_count": slide_count,
            "table_count": table_count,
            "image_count": image_count
        }

        return {
            "metadata": metadata,
            "data": data
        }

    except Exception as e:

        raise ValueError(
            f"Could not parse PPTX: {e}"
        )
    