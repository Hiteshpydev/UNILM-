from docx import Document
from pathlib import Path


def parse_docx(file_path):
    """
    Extract paragraphs, tables, and embedded images
    from a Word document.
    """

    try:
        document = Document(file_path)

        data = []

        paragraph_count = 0
        table_count = 0
        image_count = 0

        image_dir = Path("extracted_images")
        image_dir.mkdir(exist_ok=True)

        # --------------------------------
        # 1. Paragraphs
        # --------------------------------

        for paragraph_number, paragraph in enumerate(
            document.paragraphs,
            start=1
        ):

            text = paragraph.text.strip()

            if not text:
                continue

            paragraph_count += 1

            data.append({
                "location": {
                    "type": "paragraph",
                    "number": paragraph_number
                },
                "content": {
                    "type": "text",
                    "text": text
                }
            })

        # --------------------------------
        # 2. Tables
        # --------------------------------

        for table_number, table in enumerate(
            document.tables,
            start=1
        ):

            table_data = []

            for row in table.rows:

                row_data = []

                for cell in row.cells:

                    row_data.append(
                        cell.text.strip()
                    )

                table_data.append(row_data)

            if not any(
                any(cell for cell in row)
                for row in table_data
            ):
                continue

            table_count += 1

            data.append({
                "location": {
                    "type": "table",
                    "number": table_number
                },
                "content": {
                    "type": "table",
                    "data": table_data
                }
            })

        # --------------------------------
        # 3. Embedded images
        # --------------------------------

        for relationship in document.part.rels.values():

            if "image" not in relationship.reltype:
                continue

            image_part = relationship.target_part

            image_count += 1

            image_ext = (
                image_part.content_type
                .split("/")[-1]
            )

            image_path = (
                image_dir
                / f"{Path(file_path).stem}"
                f"_image_{image_count}"
                f".{image_ext}"
            )

            with open(
                image_path,
                "wb"
            ) as image_file:

                image_file.write(
                    image_part.blob
                )

            data.append({
                "location": {
                    "type": "image",
                    "number": image_count
                },
                "content": {
                    "type": "image",
                    "image_path": str(
                        image_path
                    ),
                    "source": "embedded"
                }
            })

        # --------------------------------
        # 4. Metadata
        # --------------------------------

        metadata = {
            "source": Path(file_path).name,
            "file_type": "docx",
            "paragraph_count": paragraph_count,
            "table_count": table_count,
            "image_count": image_count
        }

        return {
            "metadata": metadata,
            "data": data
        }

    except Exception as e:

        raise ValueError(
            f"Could not parse DOCX: {e}"
        )